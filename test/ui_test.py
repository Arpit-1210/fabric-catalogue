# End-to-end UI test against an in-memory fake of Supabase (the real database rules were tested separately in Postgres).
import json, base64, time, uuid, re, io, sys, os
from urllib.parse import urlparse, parse_qs
from playwright.sync_api import sync_playwright
from PIL import Image
import openpyxl

def jpg(color):
    b = io.BytesIO(); Image.new('RGB', (900, 700), color).save(b, 'JPEG'); return b.getvalue()
IMG = jpg((200, 80, 120))
os.makedirs('/tmp/uit', exist_ok=True)
open('/tmp/uit/a.jpg', 'wb').write(IMG); open('/tmp/uit/b.jpg', 'wb').write(jpg((40, 120, 200)))
open('/tmp/uit/bad.jpg', 'wb').write(b'not an image at all')

def tok(email):
    b = lambda o: base64.urlsafe_b64encode(json.dumps(o).encode()).rstrip(b'=').decode()
    return f"{b({'alg':'HS256'})}.{b({'sub':'1','email':email,'exp':int(time.time())+3600,'role':'authenticated'})}.sig"

now = lambda: time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime())
DB = {
 'fp_categories': [{'id': 'c1', 'name': 'Roto Print', 'prefix': 'RP', 'next_number': 1, 'active': True}, {'id': 'c2', 'name': 'Velvet Foil', 'prefix': 'VF', 'next_number': 1, 'active': True}],
 'fp_units': [{'name': n} for n in ['kg', 'meters', 'pieces', 'rolls']],
 'fp_designs': [], 'fp_shades': [], 'fp_stock_moves': [], 'fp_audit': [],
 'fp_staff': [{'email': 'admin@gfpl.com', 'name': 'Administrator', 'role': 'admin', 'active': True, 'created_at': now()}],
}
STORE = {}
USER = {'email': 'admin@gfpl.com'}
LOG = []

def match(rows, q):
    out = rows
    for k, v in q.items():
        if k in ('select', 'order', 'limit', 'offset', 'columns'): continue
        m = re.match(r'(eq|in)\.(.*)', v[0]);
        if m and m.group(1) == 'eq': out = [r for r in out if str(r.get(k)).lower() == m.group(2).lower() or str(r.get(k)) == m.group(2)]
    return out

def handle(route):
    req = route.request; u = urlparse(req.url); q = parse_qs(u.query); p = u.path; m = req.method
    H = {'access-control-allow-origin': '*', 'access-control-allow-headers': '*', 'access-control-allow-methods': '*', 'access-control-expose-headers': '*'}
    def J(d, s=200): route.fulfill(status=s, headers=H, content_type='application/json', body=json.dumps(d))
    if m == 'OPTIONS': return route.fulfill(status=204, headers=H)
    LOG.append((m, p))
    if p.endswith('/auth/v1/token'):
        body = json.loads(req.post_data or '{}')
        if body.get('password') != 'secret123': return J({'error': 'invalid_grant', 'error_description': 'Invalid login credentials', 'msg': 'Invalid login credentials'}, 400)
        USER['email'] = body['email'].lower()
        t = tok(USER['email'])
        return J({'access_token': t, 'refresh_token': 'r', 'expires_in': 3600, 'expires_at': int(time.time()) + 3600, 'token_type': 'bearer', 'user': {'id': '1', 'email': USER['email'], 'aud': 'authenticated'}})
    if p.endswith('/auth/v1/user'): return J({'id': '1', 'email': USER['email'], 'aud': 'authenticated'})
    if p.endswith('/auth/v1/logout'): return route.fulfill(status=204, headers=H)
    # storage
    if '/storage/v1/object/sign/fabrics' in p and m == 'POST':
        paths = json.loads(req.post_data)['paths']
        return J([{'path': x, 'signedURL': f'/object/sign/fabrics/{x}?token=t', 'error': None} for x in paths])
    if '/storage/v1/object/sign/fabrics/' in p and m == 'GET': return route.fulfill(status=200, headers={**H, 'content-type': 'image/jpeg'}, body=IMG)
    mm = re.search(r'/storage/v1/object/fabrics/(.+)$', p)
    if mm and m == 'POST':
        if mm.group(1) in STORE: return J({'error': 'Duplicate', 'message': 'The resource already exists'}, 400)
        raw = req.post_data_buffer; a = raw.find(b'\xff\xd8'); z = raw.rfind(b'\xff\xd9'); STORE[mm.group(1)] = raw[a:z + 2] if a >= 0 else raw; return J({'Key': 'fabrics/' + mm.group(1)})
    if mm and m == 'GET': return route.fulfill(status=200, headers={**H, 'content-type': 'image/jpeg'}, body=STORE.get(mm.group(1), IMG))
    # rpc
    mm = re.search(r'/rest/v1/rpc/(\w+)$', p)
    if mm:
        fn = mm.group(1); a = json.loads(req.post_data or '{}'); adm = USER['email'] == 'admin@gfpl.com'
        if fn == 'fp_whoami':
            s = [x for x in DB['fp_staff'] if x['email'] == USER['email'] and x['active']]
            return J({'email': s[0]['email'], 'name': s[0]['name'], 'role': s[0]['role']} if s else None)
        if fn == 'fp_create_design':
            c = [x for x in DB['fp_categories'] if x['id'] == a['p']['category_id']][0]
            code = f"{c['prefix']}-{c['next_number']:03d}"; c['next_number'] += 1
            d = {'id': str(uuid.uuid4()), 'code': code, 'category_id': c['id'], 'name': a['p']['name'], 'colours': a['p'].get('colours', ''), 'notes': a['p'].get('notes', ''), 'fabric_width': a['p'].get('fabric_width', ''), 'image_path': a['p'].get('image_path'), 'thumb_path': a['p'].get('thumb_path'), 'archived': False, 'created_by': USER['email'], 'updated_by': USER['email'], 'created_at': now(), 'updated_at': now()}
            DB['fp_designs'].append(d); return J(d)
        if fn == 'fp_stock_apply':
            s = [x for x in DB['fp_shades'] if x['id'] == a['p_shade']][0]
            new = a['p_qty'] if a['p_mode'] == 'set' else float(s['stock']) + a['p_qty']
            if new < 0: return J({'message': f"Not enough stock: {s['shade_code']} has {s['stock']} {s['unit']}"}, 400)
            DB['fp_stock_moves'].append({'id': str(uuid.uuid4()), 'shade_id': s['id'], 'delta': new - float(s['stock']), 'balance': new, 'kind': a['p_mode'], 'note': a['p_note'], 'by_email': USER['email'], 'created_at': now()})
            s['stock'] = new; return J(s)
        if fn == 'fp_stock_bulk':
            for it in a['p_items']:
                s = [x for x in DB['fp_shades'] if x['id'] == it['shade_id']][0]
                if it.get('unit'): s['unit'] = it['unit']
                s['stock'] = it['qty'] if a['p_mode'] == 'set' else float(s['stock']) + it['qty']
            return J({'updated': len(a['p_items'])})
        if fn == 'fp_add_staff':
            DB['fp_staff'].append({'email': a['p_email'].lower(), 'name': a['p_name'], 'role': a['p_role'], 'active': True, 'created_at': now()}); return J(a['p_email'])
        return J(None)
    mm = re.search(r'/rest/v1/(\w+)$', p)
    if mm:
        t = mm.group(1); rows = DB[t]
        if m == 'GET':
            r = match(rows, q); off = int(q.get('offset', ['0'])[0]); lim = int(q.get('limit', ['100000'])[0]); r = r[off:off + lim]
            if 'pgrst.object' in (req.headers.get('accept') or ''): return J(r[0] if r else None)
            return J(r)
        body = json.loads(req.post_data or 'null')
        if m == 'POST':
            body = body if isinstance(body, list) else [body]; out = []
            for b in body:
                if t == 'fp_shades':
                    if any(x['design_id'] == b['design_id'] and x['shade_code'].lower() == b['shade_code'].lower() for x in rows): return J({'message': 'duplicate key value violates unique constraint "fp_shades_code_uq"'}, 409)
                    b = {**b, 'id': str(uuid.uuid4()), 'stock': 0, 'archived': False, 'created_at': now(), 'updated_at': now(), 'created_by': USER['email'], 'image_path': b.get('image_path'), 'thumb_path': b.get('thumb_path')}
                if t == 'fp_categories': b = {**b, 'id': str(uuid.uuid4()), 'next_number': 1, 'active': True}
                rows.append(b); out.append(b)
            return J(out if 'pgrst.object' not in (req.headers.get('accept') or '') else out[0], 201)
        if m == 'PATCH':
            r = match(rows, q)
            for x in r: x.update(body); x['updated_at'] = now()
            return J(r if 'pgrst.object' not in (req.headers.get('accept') or '') else (r[0] if r else None))
    return J({'message': 'unhandled ' + p}, 404)

def run():
    res = []
    def ok(name, cond, extra=''):
        res.append((name, bool(cond))); print(('PASS ' if cond else 'FAIL ') + name + (' — ' + str(extra) if extra and not cond else ''))
    with sync_playwright() as pw:
        br = pw.chromium.launch(); ctx = br.new_context(viewport={'width': 390, 'height': 800}, accept_downloads=True); pg = ctx.new_page()
        errs = []; pg.on('pageerror', lambda e: errs.append(str(e)))
        ctx.route(re.compile(r'.*supabase\.co.*'), handle)
        pg.goto('http://localhost:4173/'); pg.wait_for_selector('#lf')
        pg.fill('#le', 'admin@gfpl.com'); pg.fill('#lp', 'wrong'); pg.click('#lg'); pg.wait_for_timeout(500)
        ok('wrong password shows friendly error', 'Wrong email or password' in pg.inner_text('#lerr'))
        pg.fill('#lp', 'secret123'); pg.click('#lg'); pg.wait_for_selector('.nav', timeout=8000)
        ok('login works; empty catalogue message', 'No designs yet' in pg.inner_text('#view'))
        ok('admin sees Admin tab', pg.locator('.nav a[data-n=admin]').count() == 1)

        pg.click('.nav a[data-n=new]'); pg.wait_for_selector('#e-name')
        pg.set_input_files('#e-im', '/tmp/uit/bad.jpg'); pg.wait_for_timeout(500)
        ok('bad photo rejected with message', 'could not be read' in pg.inner_text('#ee'), pg.inner_text('#ee'))
        pg.click('#sv'); ok('name required', 'name is required' in pg.inner_text('#ee').lower())
        pg.fill('#e-name', 'Rose Garden'); pg.fill('#e-col', 'pink, red'); pg.fill('#e-w', '58 inch'); pg.fill('#e-no', 'Big floral')
        pg.set_input_files('#e-im', '/tmp/uit/a.jpg'); pg.wait_for_selector('#pv img'); pg.click('#sv'); pg.wait_for_selector('#addshade', timeout=8000)
        ok('design created with code RP-001', 'RP-001' in pg.inner_text('#view'))
        ok('both photo copies uploaded', sum(1 for k in STORE if k.startswith('designs/')) == 2, list(STORE))
        ok('main photo displayed', pg.locator('#hero').count() == 1)

        pg.click('#addshade'); pg.fill('#f-sc', 'S-01'); pg.fill('#f-co', 'Rose pink'); pg.fill('#f-st', '100'); pg.set_input_files('#f-im', '/tmp/uit/b.jpg'); pg.wait_for_selector('.modal #pv img'); pg.click('.modal #ok'); pg.wait_for_selector('.shade', timeout=8000)
        ok('shade added with opening stock 100', '100' in pg.inner_text('.shade .qty'), pg.inner_text('.shade .qty'))
        pg.screenshot(path='/tmp/uit/s_design.png', full_page=True)
        pg.click('#addshade'); pg.fill('#f-sc', 's-01'); pg.click('.modal #ok'); pg.wait_for_timeout(500)
        ok('duplicate shade code message', 'already has a shade' in pg.inner_text('.modal #fe'), pg.inner_text('.modal #fe'))
        pg.click('.modal [data-close]')
        pg.click('[data-act=in]'); pg.fill('#sq', '20'); pg.click('.modal #ok'); pg.wait_for_timeout(400)
        pg.click('[data-act=out]'); pg.fill('#sq', '5'); pg.fill('#sn', 'sold to X'); pg.click('.modal #ok'); pg.wait_for_timeout(400)
        ok('stock now 115', '115' in pg.inner_text('.shade .qty'), pg.inner_text('.shade .qty'))
        pg.click('[data-act=out]'); pg.fill('#sq', '999'); pg.click('.modal #ok'); pg.wait_for_timeout(400)
        ok('overdraw blocked with message', 'Not enough stock' in pg.inner_text('.modal #se'), pg.inner_text('.modal #se'))
        pg.click('.modal [data-close]')
        pg.click('[data-act=hist]'); pg.wait_for_selector('.modal table'); ok('history shows movements', 'sold to X' in pg.inner_text('.modal table')); pg.click('.modal [data-close]')

        # second design in another category
        pg.click('.nav a[data-n=new]'); pg.wait_for_selector('#e-cat'); pg.select_option('#e-cat', 'c2'); pg.fill('#e-name', 'Gold Leaf'); pg.click('#sv'); pg.wait_for_selector('#addshade')
        ok('VF-001 issued for other category', 'VF-001' in pg.inner_text('#view'))
        pg.click('.nav a[data-n=new]'); pg.wait_for_selector('#e-cat'); pg.fill('#e-name', 'Paisley'); pg.click('#sv'); pg.wait_for_selector('#addshade')
        ok('sequential RP-002', 'RP-002' in pg.inner_text('#view'))

        pg.click('.nav a[data-n=home]'); pg.wait_for_selector('#q'); pg.fill('#q', 'rp1'); pg.wait_for_timeout(300)
        t = pg.inner_text('#res'); ok('search "rp1" finds RP-001 with stock & shade', 'RP-001' in t and '115' in t and 'S-01' in t and 'RP-002' not in t, t[:200])
        pg.fill('#q', 'pink'); pg.wait_for_timeout(300); ok('search by colour', 'RP-001' in pg.inner_text('#res'))
        pg.fill('#q', 'roto'); pg.wait_for_timeout(300); ok('search by category', 'RP-001' in pg.inner_text('#res') and 'VF-001' not in pg.inner_text('#res'))
        pg.fill('#q', ''); pg.wait_for_timeout(600); pg.screenshot(path='/tmp/uit/s_grid.png'); ok('grid shows 3 designs', pg.locator('.tile').count() == 3)
        pg.select_option('#fs', 'ok'); pg.wait_for_timeout(300); ok('filter in stock → 1', pg.locator('.tile').count() == 1); pg.select_option('#fs', '')

        # archive
        pg.fill('#q', 'paisley'); pg.wait_for_timeout(300); pg.click('.hit' if pg.locator('.hit').count() else '.tile'); pg.wait_for_selector('#arch'); pg.click('#arch'); pg.click('.modal #c-yes'); pg.wait_for_timeout(500)
        pg.click('.nav a[data-n=home]'); pg.wait_for_selector('#q'); pg.fill('#q', ''); pg.wait_for_timeout(300)
        ok('archived design hidden from catalogue', pg.locator('.tile').count() == 2)
        pg.select_option('#fa', '1'); pg.wait_for_timeout(300); ok('archived filter shows it', pg.locator('.tile').count() == 1 and 'RP-002' in pg.inner_text('#res'))
        pg.select_option('#fa', '0')

        # exports
        pg.click('.nav a[data-n=export]'); pg.wait_for_selector('#b-xlsx'); pg.select_option('#x-scope', 'category'); pg.wait_for_timeout(200)
        ok('export count 2', '2 designs' in pg.inner_text('#x-count'), pg.inner_text('#x-count'))
        with pg.expect_download(timeout=15000) as dl: pg.click('#b-xlsx')
        f = dl.value; path = '/tmp/uit/' + f.suggested_filename; f.save_as(path)
        wb = openpyxl.load_workbook(path); ws = wb.active; rows = list(ws.iter_rows(values_only=True))
        ok('xlsx has header + rows incl. stock', rows[0][0] == 'Print Code' and any(r[0] == 'RP-001' and r[7] == 'S-01' and r[11] == 115 for r in rows[1:]), rows[:3])
        with pg.expect_download(timeout=15000) as dl: pg.click('#b-csv')
        f = dl.value; f.save_as('/tmp/uit/x.csv'); txt = open('/tmp/uit/x.csv', encoding='utf-8-sig').read()
        ok('csv ok', 'Print Code' in txt and 'RP-001' in txt and 'S-01' in txt)
        with pg.expect_download(timeout=30000) as dl: pg.click('#b-pdf')
        f = dl.value; f.save_as('/tmp/uit/c.pdf'); b = open('/tmp/uit/c.pdf', 'rb').read(); ok('pdf grid ok', b[:4] == b'%PDF' and len(b) > 5000, len(b))
        with pg.expect_download(timeout=30000) as dl: pg.click('#b-pdfd')
        f = dl.value; f.save_as('/tmp/uit/d.pdf'); b = open('/tmp/uit/d.pdf', 'rb').read(); ok('pdf detailed ok', b[:4] == b'%PDF' and len(b) > 5000, len(b))
        pg.select_option('#x-scope', 'pick'); pg.wait_for_selector('#x-list label'); pg.fill('#x-pq', 'gold'); pg.wait_for_timeout(200); pg.click('#x-list input')
        ok('pick designs → 1 selected', '1 design ' in pg.inner_text('#x-count'), pg.inner_text('#x-count'))

        # stock via excel
        pg.click('.nav a[data-n=stock]'); pg.wait_for_selector('#tpl')
        with pg.expect_download() as dl: pg.click('#tpl')
        dl.value.save_as('/tmp/uit/tpl.xlsx'); wb = openpyxl.load_workbook('/tmp/uit/tpl.xlsx'); ws = wb.active
        ok('template lists shade with current stock', ws['A2'].value == 'RP-001' and ws['B2'].value == 'S-01' and ws['F2'].value == 115, [c.value for c in ws[2]])
        ws['G2'] = 300; ws['E2'] = 'rolls'; ws.append(['XX-999', 'Z', None, None, None, None, 5]); wb.save('/tmp/uit/fill.xlsx')
        pg.set_input_files('#xf', '/tmp/uit/fill.xlsx'); pg.wait_for_selector('#apply'); t = pg.inner_text('#pv')
        ok('preview: 1 ready, 1 skipped (unknown code)', '1\nready' in t.replace(' ', '\n') or ('ready' in t and 'skipped' in t and 'not found' in t.lower()), t[:200])
        pg.click('#apply'); pg.click('.modal #c-yes'); pg.wait_for_selector('#pv .card b', timeout=8000); pg.wait_for_timeout(500)
        sh = DB['fp_shades'][0]; ok('excel import applied (300 rolls)', sh['stock'] == 300 and sh['unit'] == 'rolls', sh)

        # admin
        pg.click('.nav a[data-n=admin]'); pg.wait_for_selector('#s-add'); pg.fill('#s-n', 'Ravi'); pg.fill('#s-e', 'ravi@shop.com'); pg.click('#s-add'); pg.wait_for_selector('.modal', timeout=5000)
        ok('staff created, credentials shown', 'ravi@shop.com' in pg.inner_text('.modal')); pg.click('.modal [data-close]')
        pg.click('.tabs [data-t=cats]'); pg.wait_for_selector('#c-add'); pg.fill('#c-n', 'Block Print'); pg.fill('#c-p', 'bp'); pg.click('#c-add'); pg.wait_for_timeout(600)
        ok('category added', any(c['prefix'] == 'BP' for c in DB['fp_categories']))
        pg.click('.tabs [data-t=backup]'); pg.wait_for_selector('#bk-j')
        with pg.expect_download() as dl: pg.click('#bk-j')
        ok('json backup', b'RP-001' in open(dl.value.path(), 'rb').read())
        with pg.expect_download(timeout=30000) as dl: pg.click('#bk-z')
        import zipfile; z = zipfile.ZipFile(dl.value.path()); names = z.namelist(); ok('zip has data.json + photos', 'data.json' in names and any(n.startswith('photos/designs/') for n in names), names[:5])

        # staff user sees no admin tab; unknown login blocked
        pg.click('#lo'); pg.wait_for_selector('#lf'); USER_BAK = USER['email']
        pg.fill('#le', 'ravi@shop.com'); pg.fill('#lp', 'secret123'); pg.click('#lg'); pg.wait_for_selector('.nav'); ok('staff has no Admin tab', pg.locator('.nav a[data-n=admin]').count() == 0)
        pg.goto('http://localhost:4173/#/admin'); pg.wait_for_timeout(1500); ok('staff cannot open admin page', '#/admin' not in pg.url)
        pg.click('#lo'); pg.wait_for_selector('#lf'); pg.fill('#le', 'pankaj@gfpl.com'); pg.fill('#lp', 'secret123'); pg.click('#lg'); pg.wait_for_selector('#lerr:not(:empty)', timeout=5000)
        ok('login from another app is refused', 'does not have access' in pg.inner_text('#lerr'), pg.inner_text('#lerr'))
        ok('no JS errors', not errs, errs)
        br.close()
    bad = [n for n, c in res if not c]; print('\n%d/%d passed' % (len(res) - len(bad), len(res))); return 1 if bad else 0

sys.exit(run())
